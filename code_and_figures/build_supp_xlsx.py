"""Builds Supplementary_Data.xlsx: one sheet per dataset, counts as values, proportions and Wilson CIs as live formulas."""
import os, pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
D = lambda f: os.path.join(HERE, 'data', f)
wb = Workbook()
ARIAL = Font(name='Arial', size=10); BOLD = Font(name='Arial', size=10, bold=True); BLUE = Font(name='Arial', size=10, color='0000FF')
HEAD = PatternFill('solid', fgColor='EEEEEE')

def write_sheet(name, df, source, note=None, k_col=None, n_col=None, derive=True):
    ws = wb.create_sheet(name[:31])
    ws['A1'] = name; ws['A1'].font = Font(name='Arial', size=12, bold=True)
    ws['A2'] = 'Source: ' + source; ws['A2'].font = Font(name='Arial', size=9, italic=True)
    r0 = 4
    if note:
        ws['A3'] = note; ws['A3'].font = Font(name='Arial', size=9, italic=True)
    cols = list(df.columns)
    extra = []
    if derive and k_col and n_col:
        extra = ['proportion', 'wilson_lo', 'wilson_hi']
    for j, c in enumerate(cols + extra, start=1):
        cell = ws.cell(row=r0, column=j, value=c); cell.font = BOLD; cell.fill = HEAD
    for i, row in enumerate(df.itertuples(index=False), start=r0 + 1):
        for j, v in enumerate(row, start=1):
            cell = ws.cell(row=i, column=j, value=(None if pd.isna(v) else v)); cell.font = BLUE if cols[j - 1] in (k_col, n_col) else ARIAL
        if extra:
            kc, nc = get_column_letter(cols.index(k_col) + 1), get_column_letter(cols.index(n_col) + 1)
            k, n = f'{kc}{i}', f'{nc}{i}'
            pcol = len(cols) + 1
            ws.cell(row=i, column=pcol, value=f'=IF({n}>0,{k}/{n},"")').font = ARIAL
            # Wilson score interval, z = 1.959964
            z = '1.959964'
            base = f'(({k}/{n})+{z}^2/(2*{n}))/(1+{z}^2/{n})'
            half = f'({z}*SQRT(({k}/{n})*(1-{k}/{n})/{n}+{z}^2/(4*{n}^2)))/(1+{z}^2/{n})'
            ws.cell(row=i, column=pcol + 1, value=f'=IF({n}>0,MAX(0,{base}-{half}),"")').font = ARIAL
            ws.cell(row=i, column=pcol + 2, value=f'=IF({n}>0,MIN(1,{base}+{half}),"")').font = ARIAL
            for cc in (pcol, pcol + 1, pcol + 2): ws.cell(row=i, column=cc).number_format = '0.000'
    for j, c in enumerate(cols + extra, start=1):
        width = max([len(str(c))] + [len(str(v)) for v in df.iloc[:, j - 1]] if j <= len(cols) else [10]) + 2
        ws.column_dimensions[get_column_letter(j)].width = min(max(width, 10), 60)
    ws.freeze_panes = ws.cell(row=r0 + 1, column=1)
    return ws

# README
ws = wb.active; ws.title = 'README'
lines = [
    ('Supplementary Data', Font(name='Arial', size=14, bold=True)),
    ('HGVS nomenclature compliance in clinical variant literature, 2010–2026: meta-analysis and interrupted time series', ARIAL),
    ('Prepared 22 October 2026. Companion to the main manuscript, Supplementary Information, analysis.py and analysis_supplementary.py.', ARIAL),
    ('', ARIAL),
    ('Conventions', BOLD),
    ('Blue values are counts or denominators taken from the cited paper (or back-computed from its percentage and denominator, see each sheet note).', ARIAL),
    ('Black values in the proportion / wilson_lo / wilson_hi columns are formulas (Wilson score 95% interval, z = 1.959964) and recalculate if a count is changed.', ARIAL),
    ('Percent columns copied from a paper are stored as reported (e.g., 12.1 means 12.1%); proportion columns are fractions.', ARIAL),
    ('', ARIAL),
    ('Sheets', BOLD),
    ('studies — study-level estimates entering Figure 1 (forest plot)', ARIAL),
    ('eqa_yearly — total EQA series by scheme year (interrupted time series, Figure 2)', ARIAL),
    ('eqa_provider_yearly — the four provider–indication sub-series (Figure 3, Figure S1)', ARIAL),
    ('eqa_participations — compliance by number of successive EQA participations (Figure S3)', ARIAL),
    ('eqa_transitions — compliance conditional on the previous scheme result', ARIAL),
    ('eqa_methods — compliance by test method', ARIAL),
    ('eqa_genes — compliance and serious errors by gene', ARIAL),
    ('eqa_error_types — composition of EQA outcomes', ARIAL),
    ('eqa_returning_variants — variants distributed in more than one scheme year (Figure S4)', ARIAL),
    ('manuscript_criteria — rubric results for 52 Genetics in Medicine / GIM Open manuscripts (Figure S5)', ARIAL),
    ('findability — 18 published variants searched in three tools (Figure S6)', ARIAL),
    ('tool_agreement — agreement of HGVS descriptions between annotation tools and ClinVar', ARIAL),
    ('journal_pilot — HUGO pilot of validation-file requirements at two journals', ARIAL),
    ('mutalyzer_assessments — Name Checker outcome composition', ARIAL),
    ('hgvs_versions — complete HGVS Nomenclature release log', ARIAL),
    ('its_series — observed, fitted and counterfactual values of the main ITS model', ARIAL),
]
for i, (t, f) in enumerate(lines, start=1):
    ws.cell(row=i, column=1, value=t).font = f
ws.column_dimensions['A'].width = 120

K = 'https://pmc.ncbi.nlm.nih.gov/articles/PMC6973115'; L = 'https://academic.oup.com/clinchem/article/72/6/652/8510568'
st = pd.read_csv(os.path.join(HERE, '..', 'studies.csv'))
write_sheet('studies', st, 'as listed in the source column of each row', 'events back-computed from the published percentage × n where the paper reports only a percentage (Keppens 490; Lefter 13.1M; Chen 96,294 and 138,287; Berwouts 139)', 'events', 'n')
ey = pd.read_csv(os.path.join(HERE, '..', 'eqa_yearly.csv'))
ey['nomen_k'] = (ey.nomen_pct / 100 * ey.nomen_n).round().astype(int); ey['refseq_k'] = (ey.refseq_pct / 100 * ey.refseq_n).round().astype(int)
write_sheet('eqa_yearly', ey[['year', 'nomen_pct', 'nomen_k', 'nomen_n', 'refseq_pct', 'refseq_k', 'refseq_n']], 'Keppens et al. 2020, Hum Mutat 41:7, Table 3 (Total column); ' + K, 'nomen_k and refseq_k back-computed as round(pct × n / 100)', 'nomen_k', 'nomen_n')
pv = pd.read_csv(D('eqa_provider_yearly.csv')); pv['nomen_k'] = (pv.nomen_pct / 100 * pv.nomen_n).round().astype(int); pv['refseq_k'] = (pv.refseq_pct / 100 * pv.refseq_n).round().astype(int)
write_sheet('eqa_provider_yearly', pv[['provider', 'indication', 'year', 'nomen_pct', 'nomen_k', 'nomen_n', 'refseq_pct', 'refseq_k', 'refseq_n']], 'Keppens et al. 2020, Table 3 (provider columns); ' + K, 'k columns back-computed as round(pct × n / 100); no ESP mCRC scheme in 2015; Gen&Tiss series end in 2017', 'nomen_k', 'nomen_n')
pp = pd.read_csv(D('eqa_participations.csv')); pp['nomen_k'] = (pp.nomen_pct / 100 * pp.nomen_n).round().astype(int); pp['refseq_k'] = (pp.refseq_pct / 100 * pp.refseq_n).round().astype(int)
write_sheet('eqa_participations', pp[['participations', 'nomen_pct', 'nomen_k', 'nomen_n', 'refseq_pct', 'refseq_k', 'refseq_n']], 'Keppens et al. 2020, Table 3 (number of EQA participations); ' + K, 'k columns back-computed', 'nomen_k', 'nomen_n')
tr = pd.read_csv(D('eqa_transitions.csv')); tr['k'] = (tr.pct_compliant / 100 * tr.n).round().astype(int)
write_sheet('eqa_transitions', tr[['previous_result', 'outcome', 'pct_compliant', 'k', 'n']], 'Keppens et al. 2020, Table 3 (performance between two participations); ' + K, 'k back-computed', 'k', 'n')
me = pd.read_csv(D('eqa_methods.csv')); me['k'] = (me.pct_compliant / 100 * me.n).round().astype(int)
write_sheet('eqa_methods', me[['method', 'outcome', 'pct_compliant', 'k', 'n']], 'Keppens et al. 2020, Figure 1 and text; ' + K, 'k back-computed; overall NGS compliance in the text is 20.9% of 447', 'k', 'n')
ge = pd.read_csv(D('eqa_genes.csv')); ge['k_compliant'] = (ge.pct_compliant / 100 * ge.n).round().astype(int); ge['k_serious'] = (ge.pct_serious_error / 100 * ge.n_serious).round().astype(int)
write_sheet('eqa_genes', ge[['gene', 'pct_compliant', 'k_compliant', 'n', 'pct_serious_error', 'k_serious', 'n_serious']], 'Keppens et al. 2020, text (Table S1 of that paper); ' + K, 'k columns back-computed', 'k_compliant', 'n')
write_sheet('eqa_error_types', pd.read_csv(D('eqa_error_types.csv')), 'Keppens et al. 2020, Results 3.1; ' + K, None, None, None, derive=False)
rv = pd.read_csv(D('eqa_returning_variants.csv')); rv['k'] = (rv.pct / 100 * rv.n).round().astype(int)
write_sheet('eqa_returning_variants', rv[['gene', 'variant', 'year', 'pct', 'k', 'n']], 'Keppens et al. 2020, Table 2; ' + K, 'k back-computed; averaged percentages where a variant ran in both schemes in a year, as presented by the paper', 'k', 'n')
write_sheet('manuscript_criteria', pd.read_csv(D('manuscript_criteria.csv')), 'Lansdon et al. 2026, Clin Chem 72:652, Table 1; ' + L, 'n = critical + small + none for each criterion', 'any_error', 'n')
write_sheet('findability', pd.read_csv(D('findability.csv')), 'Lansdon et al. 2026, Results (Assessment of variant nomenclature); ' + L, None, 'found', 'n')
write_sheet('tool_agreement', pd.read_csv(D('tool_agreement.csv')), 'Yen et al. 2017 Genome Med 9:7 (PMC5267466); Chen et al. 2025 Hum Genomics 19:70 (bioRxiv 2024.10.11.617756)', 'Chen et al. counts back-computed from match fractions × 164,549', 'agree', 'n')
jp = pd.read_csv(D('journal_pilot.csv'))
write_sheet('journal_pilot', jp, 'Higgins et al. 2021, Hum Mutat 42:3 (accepted manuscript, Cardiff ORCA)', None, 'with_validation_file', 'with_variants')
write_sheet('mutalyzer_assessments', pd.read_csv(D('mutalyzer_assessments.csv')), 'Lefter et al. 2021, Bioinformatics 37:2811, Figure 3 and Section 6.2', None, None, None, derive=False)
write_sheet('hgvs_versions', pd.read_csv(D('hgvs_versions.csv')), 'hgvs-nomenclature.org version pages (accessed 20 Oct 2026); den Dunnen et al. 2016', None, None, None, derive=False)
its = pd.read_csv(os.path.join(HERE, '..', 'its_data.csv'))
write_sheet('its_series', its, 'computed by analysis.py (segmented binomial model, break 2016)', 'fitted = model with level and slope change at 2016; counterfactual = pre-2016 model projected forward', 'nomen_k', 'nomen_n')

out = os.path.join(HERE, 'out', 'Supplementary_Data.xlsx'); wb.save(out); print('written', out)
